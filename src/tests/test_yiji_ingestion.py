"""Part-four ingestion, namespace, provenance, and retrieval regressions."""

import importlib
import json
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

import database as db
import retrieval
import chat_engine as ce
from evidence import EvidenceRegistry
from tools import ingest_yiji_textbook as ingest


@pytest.fixture
def entries():
    return ingest.parse_docx(ingest.DEFAULT_DOCX)


@pytest.fixture
def yiji_db(tmp_path, monkeypatch, entries):
    monkeypatch.setattr(db, "_connection", None)
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "content.db"))
    db.init_db()
    db.save_fuling_article(2, "普通第二条", 83, "普通宋本对照", "taiyang")
    db.save_zabing_article("zabing_26.9", "26.9", "杂病已有条文", "10.9", "对照", "金匮", "第二十六章", "part3.docx")
    db.save_conversation("preserve", "test@example.invalid", "2026-01-01", [{"role": "user", "content": "retain me"}])
    db.save_lesson("lesson0095", "Supplement", "additional", "original", "lecture.docx", "retain this lecture")
    ingest.load_entries(entries)
    yield
    db.close_db()


def test_source_inventory_and_numbering_gap(entries, caplog):
    assert len(entries) == 113
    assert len({e["chapter_title"] for e in entries}) == 15
    assert [e["fuling_ref"] for e in entries] == [f"YJ.{n}" for n in range(1, 115) if n != 31]
    assert sum(bool(e["comparison_ref"]) for e in entries) == 5
    assert all(e["source_path"] == ingest.DEFAULT_DOCX.name for e in entries)
    ingest.parse_docx(ingest.DEFAULT_DOCX)
    assert "YJ.31" in caplog.text


def test_original_comparison_and_section_are_separate(entries):
    second = entries[1]
    assert second["fuling_zh"] == "咽喉干燥者，忌发其汗。"
    assert second["comparison_ref"] == "83"
    assert second["comparison_zh"] == "咽喉干燥者，不可发汗。"
    assert second["comparison_book"] == "宋本"
    assert second["chapter_title"].endswith("忌发汗第一")
    assert entries[0]["comparison_ref"] is None
    assert "忌发汗第一" not in entries[0]["fuling_zh"]
    assert entries[-1]["chapter_title"].endswith("宜水第十五")
    assert "五苓散" in entries[-1]["fuling_zh"]


def test_duplicate_and_unassigned_content_fail_before_import(tmp_path):
    path = tmp_path / "bad.docx"
    doc = Document()
    for text in ("辨伤寒宜忌脉症篇第十五", "忌发汗第一", "涪陵古本 宜忌YJ.1. 【原文】", "原文", "涪陵古本 宜忌YJ.1. 【原文】", "重复"):
        doc.add_paragraph(text)
    doc.save(path)
    with pytest.raises(ValueError, match="unique"):
        ingest.parse_docx(path)
    doc = Document()
    doc.add_paragraph("Unexpected preface")
    doc.save(path)
    with pytest.raises(ValueError, match="outside"):
        ingest.parse_docx(path)


def test_export_matches_checked_in_text_and_does_not_invent_alignments(entries, tmp_path):
    output = tmp_path / "textbook_yiji.txt"
    ingest.write_normalized_text(entries, output)
    text = output.read_text()
    assert text == (ROOT / "textbook_yiji.txt").read_text()
    assert "YJ.31:" not in text
    assert text.count("（宋本") == 5
    assert "金匮" not in text
    assert "FORMULA" not in text


def test_reimport_preserves_existing_data(yiji_db, entries):
    ingest.load_entries(entries)
    connection = db.get_connection()
    assert connection.execute("SELECT count(*) FROM zabing_articles WHERE fuling_ref LIKE 'YJ.%'").fetchone()[0] == 113
    assert db.get_zabing_by_reference("26.9")[0]["fuling_zh"] == "杂病已有条文"
    assert db.get_fuling_article(2)["fuling_zh"] == "普通第二条"
    assert db.get_messages("preserve")[0]["content"] == "retain me"
    assert db.get_lesson("lesson0095")["content"] == "retain this lecture"


def test_zabing_clear_preserves_the_guide(yiji_db, monkeypatch):
    from tools import ingest_zabing_textbooks as zabing
    existing = db.get_zabing_by_reference("26.9")[0]
    monkeypatch.setattr(zabing, "parse_docx", lambda path: ([existing], [], []))
    monkeypatch.setattr(sys, "argv", ["ingest_zabing_textbooks.py", str(ingest.DEFAULT_DOCX), "--db", "--clear"])
    zabing.main()
    assert len(db.search_yiji_articles("宜忌", search_depth="deep", limit=200)) == 113
    assert db.get_zabing_by_reference("26.9")[0]["fuling_zh"] == "杂病已有条文"


@pytest.mark.parametrize("query", ["YJ.2", "yj.2", "YJ 2", "宜忌第2条", "涪陵古本 宜忌YJ.2"])
def test_guide_references_do_not_become_normal_line_numbers(yiji_db, query):
    assert retrieval.parse_reference(query).kind == "yiji"
    records = retrieval.search_textbook(query)
    assert [r["fuling_article_num"] for r in records] == ["YJ.2"]
    assert records[0]["channel"] == "yiji"
    assert records[0]["comparison_book"] == "宋本"
    assert retrieval.search_formulas(query) == []
    assert retrieval.search_lectures(query) == []
    assert [r["fuling_article_num"] for r in retrieval.search_textbook("2")] == [2]


def test_missing_reference_is_not_filled_by_an_unrelated_line(yiji_db):
    assert retrieval.search_textbook("YJ.31") == []
    assert ce.build_context("YJ.31")[1] == []


def test_guide_keywords_and_songben_alignment(yiji_db):
    assert retrieval.search_textbook("宜忌 咽喉干燥")[0]["fuling_article_num"] == "YJ.2"
    assert retrieval.search_textbook("宜忌 五苓散")[0]["fuling_article_num"] == "YJ.114"
    assert "YJ.2" in [r["fuling_article_num"] for r in retrieval.search_textbook("咽喉干燥")]
    assert "YJ.2" in [r["fuling_article_num"] for r in retrieval.search_textbook("宋本第83条")]
    assert retrieval.search_textbook("金匮 10.9")[0]["fuling_article_num"] == "26.9"


def test_public_search_and_ai_citations_include_guide(yiji_db):
    server = importlib.import_module("src.server")
    client = server.app.test_client()
    with client.session_transaction() as session:
        session["user"] = "prof@tcm.org"
    data = client.get("/api/search?q=YJ.2&mode=name").get_json()
    assert [r["fuling_article_num"] for r in data["textbook_entries"]] == ["YJ.2"]
    registry = EvidenceRegistry()
    context, sources = ce.build_context("YJ.2", evidence=registry)
    assert sources[0]["source_id"] == "yiji_article:yiji_YJ.2"
    assert "忌发汗第一" in context
    assert "YJ.2" in sources[0]["title"] and "宋本" in sources[0]["title"]
    tool = json.loads(ce.tool_search_articles("YJ.114", evidence=registry))
    assert tool["fuling"][0]["citation"] == "[2]"
    assert registry.sources[1]["source_id"] == "yiji_article:yiji_YJ.114"
