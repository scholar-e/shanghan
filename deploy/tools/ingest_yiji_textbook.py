#!/usr/bin/env python3
"""Import textbook part 4's 宜忌 guide, preserving its YJ reference namespace."""

import argparse
import logging
import re
import sys
from pathlib import Path

from docx import Document

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR / "src"))
DEFAULT_DOCX = ROOT_DIR / "textbook4thpart_ReEdited_SHLYiJiGuide.docx"
OUT_TXT = ROOT_DIR / "textbook_yiji.txt"
ENTRY = re.compile(r"^涪陵古本\s*宜忌\s*YJ\.(\d+)\.\s*【原文】\s*(.*)$", re.I)
CHAPTER = re.compile(r"^辨.+篇第[一二三四五六七八九十百]+$")
SECTION = re.compile(r"^[宜忌].+第[一二三四五六七八九十百]+$")
COMPARISON = re.compile(r"^（宋本\s*(\d+)）\s*(.*)$")
logger = logging.getLogger(__name__)


def parse_docx(path):
    """Parse original and comparison text separately; reject unassigned text."""
    document = Document(path)
    if document.tables:
        raise ValueError("Unexpected tables in Yiji source; inspect before importing")
    chapter = section = ""
    entries = []
    current = None
    reading_comparison = False

    def finish():
        if current is None:
            return
        current["fuling_zh"] = " ".join(current["fuling_zh"])
        current["comparison_zh"] = " ".join(current["comparison_zh"]) or None
        if not current["fuling_zh"]:
            raise ValueError(f"Empty original text for {current['fuling_ref']}")
        entries.append(current)

    for paragraph in document.paragraphs:
        line = re.sub(r"\s+", " ", paragraph.text).strip()
        if not line:
            continue
        if CHAPTER.fullmatch(line) or SECTION.fullmatch(line):
            finish()
            current = None
            if CHAPTER.fullmatch(line):
                chapter, section = line, ""
            else:
                section = line
            continue
        match = ENTRY.fullmatch(line)
        if match:
            finish()
            if not chapter or not section:
                raise ValueError(f"Missing chapter/section for {line}")
            reference = f"YJ.{int(match[1])}"
            current = {
                "entry_key": f"yiji_{reference}", "fuling_ref": reference,
                "fuling_zh": [match[2]] if match[2] else [],
                "comparison_ref": None, "comparison_zh": [], "comparison_book": "宋本",
                "chapter_title": f"{chapter} / {section}", "source_path": Path(path).name,
            }
            reading_comparison = False
            continue
        if current is None:
            raise ValueError(f"Text outside a Yiji entry: {line[:80]}")
        comparison = COMPARISON.fullmatch(line)
        if comparison:
            if current["comparison_ref"] is not None:
                raise ValueError(f"Multiple comparison references for {current['fuling_ref']}")
            current["comparison_ref"] = comparison[1]
            current["comparison_zh"].append(comparison[2])
            reading_comparison = True
        else:
            if line.startswith("涪陵古本") or line.startswith("（宋本"):
                raise ValueError(f"Unrecognized reference: {line}")
            current["comparison_zh" if reading_comparison else "fuling_zh"].append(line)
    finish()
    numbers = [int(entry["fuling_ref"].split(".")[1]) for entry in entries]
    if not numbers or numbers[0] != 1 or numbers != sorted(set(numbers)):
        raise ValueError("Yiji references must be unique, increasing, and start at YJ.1")
    missing = sorted(set(range(1, numbers[-1] + 1)) - set(numbers))
    if missing:
        logger.warning("Source numbering gaps preserved: %s", ", ".join(f"YJ.{n}" for n in missing))
    return entries


def write_normalized_text(entries, path=OUT_TXT):
    lines = []
    previous_section = None
    for entry in entries:
        if entry["chapter_title"] != previous_section:
            lines.extend(["", entry["chapter_title"]])
            previous_section = entry["chapter_title"]
        lines.append(f"涪陵古本 宜忌 {entry['fuling_ref']}: {entry['fuling_zh']}")
        if entry["comparison_ref"]:
            lines.append(f"（宋本 {entry['comparison_ref']}）{entry['comparison_zh']}")
    Path(path).write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def load_entries(entries, database_path=None):
    import database as db

    if database_path:
        db.close_db()
        db.DB_PATH = str(Path(database_path).resolve())
    db.init_db()
    # One transaction, and no deletion or replacement of conversations/lectures.
    with db.get_connection() as connection:
        connection.executemany(
            """INSERT INTO zabing_articles
               (entry_key, fuling_ref, fuling_zh, comparison_ref, comparison_zh,
                comparison_book, chapter_title, source_path)
               VALUES (:entry_key, :fuling_ref, :fuling_zh, :comparison_ref, :comparison_zh,
                       :comparison_book, :chapter_title, :source_path)
               ON CONFLICT(entry_key) DO UPDATE SET
                   fuling_ref=excluded.fuling_ref, fuling_zh=excluded.fuling_zh,
                   comparison_ref=excluded.comparison_ref, comparison_zh=excluded.comparison_zh,
                   comparison_book=excluded.comparison_book, chapter_title=excluded.chapter_title,
                   source_path=excluded.source_path""", entries,
        )
    logger.info("Imported %d Yiji entries into %s", len(entries), db.DB_PATH)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", nargs="?", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--txt", action="store_true", help="Write textbook_yiji.txt")
    parser.add_argument("--db", action="store_true", help="Upsert Yiji entries into SQLite")
    parser.add_argument("--database", type=Path, help="Override the SQLite destination")
    args = parser.parse_args()
    log_dir = ROOT_DIR / "logs"
    log_dir.mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(filename)s:%(lineno)d %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(sys.stdout), logging.FileHandler(log_dir / "ingest_yiji.log")])
    entries = parse_docx(args.docx)
    logger.info("Parsed %d entries, %d sections, %d Songben alignments from %s", len(entries),
                len({e["chapter_title"] for e in entries}), sum(bool(e["comparison_ref"]) for e in entries), args.docx)
    if args.txt:
        write_normalized_text(entries)
        logger.info("Wrote %s", OUT_TXT)
    if args.db:
        load_entries(entries, args.database)


if __name__ == "__main__":
    main()
